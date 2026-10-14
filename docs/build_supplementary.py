"""Build the Supplementary Material (docs/supplementary_material.docx) from the results files.

Requires python-docx; convert to PDF with LibreOffice:
    python docs/build_supplementary.py
    soffice --headless --convert-to pdf --outdir docs docs/supplementary_material.docx
"""

import json, numpy as np
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

R = "results/"
ACC = json.load(open(R + "accuracy_table.json"))
A = json.load(open(R + "analysis.json"))
D = json.load(open(R + "decision_nested.json"))
DG = json.load(open(R + "diagnostics.json"))
SP = json.load(open(R + "supplementary_analyses.json"))
AB = json.load(open(R + "cnn_floor_summary.json"))
P = {
    "nested": D,
    "P1": json.load(open(R + "decision_P1.json")),
    "P2": json.load(open(R + "decision_P2.json")),
}
SUBS = ["FD001", "FD002", "FD003", "FD004"]
neg = lambda s: s.replace("-", "−")


def f1(x):
    return neg(f"{x:.1f}")


doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
for a in ("left_margin", "right_margin"):
    setattr(sec, a, Mm(20))
sec.top_margin = sec.bottom_margin = Mm(20)


def H(t, lvl=1):
    p = doc.add_heading(t, level=lvl)
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.color.rgb = None
        r.font.size = Pt(14 if lvl == 0 else 12 if lvl == 1 else 11)


def Pp(t, it=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.italic = it
    if size:
        r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(5)
    return p


def cap(label, t):
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    r.bold = True
    r.font.size = Pt(9.5)
    r2 = p.add_run(t)
    r2.font.size = Pt(9.5)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)


def border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for e, v in kw.items():
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(v))
        b.append(el)
    tcPr.append(b)


def table(rows, widths, size=8.5, groups=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.width = Mm(widths[j])
            c.text = ""
            p = c.paragraphs[0]
            r = p.add_run(str(v))
            r.font.size = Pt(size)
            r.bold = i == 0
            p.paragraph_format.space_after = Pt(0)
            if i == 0:
                border(c, top=8, bottom=6)
            if i == 0 and j == 0:
                trPr = t.rows[0]._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            if i == len(rows) - 1:
                border(c, bottom=8)
            if groups and i in groups and i != 0:
                border(c, top=3)
    t.autofit = False
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j] * 56.7)))
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def fig(path, w, label, t):
    doc.add_picture(path, width=Mm(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    r.bold = True
    r.font.size = Pt(9.5)
    r2 = p.add_run(t)
    r2.font.size = Pt(9.5)


# ---------------- title
H("Supplementary Material", 0)
Pp(
    "Lowest Error, Best Decision? Accuracy Equivalence and Lead-Time-Dependent Rankings of Remaining Useful Life Models for Turbofan Maintenance",
    it=True,
)
Pp(
    'Hashem Mohammad and Ali Alsayegh, University of Strathclyde. All quantities are in engine cycles unless stated. LR: linear regression; RF: random forest; GB: gradient boosting; CNN: 1D convolutional reference. "Nested" refers to the nested engine-grouped validation of Sect. 3.4; P1 and P2 are the two repeated shuffled partitions (non-nested calibration, margins re-estimated).'
)
# ---------------- S0 procedure
H("S0 Resampling and Statistical Procedure")
Pp(
    "Partition and seeds. Each subset was split into five engine-grouped folds by an unshuffled GroupKFold, which is deterministic; seeds 42 to 46 initialised the five outer fold models (random forest, gradient boosting and CNN) and seeds 142 + 10k + c the inner pair models, so they do not affect the partition. Test-set accuracy is reported over the five fold models (mean ± SD, Table S1); paired accuracy tests use the fold-averaged prediction for each test engine."
)
Pp(
    "Nested calibration. For outer fold k, models fitted on the other four folds predict the run-to-failure trajectories of fold k (outer predictions). For each calibration fold c ≠ k, models fitted on the three folds other than k and c predict fold c (inner predictions); 10 pair fits and 5 outer fits per model and subset. Preprocessing (regime clustering, scaling, variance screen) was refitted in every fit. For engine i and lead L, the required margin r(i, L) is the smallest m for which the first cycle with predicted RUL ≤ L + m has true RUL ≥ L; because the alarm time is monotone in m, the zero-miss margin for fold k is the largest r(i, L) over its calibration engines and the 2% margin is the (⌊0.02n⌋ + 1)-th largest."
)
Pp(
    "Bootstrap. In each of B = 2,000 replicates, engines were drawn with replacement within each outer fold; every fold's margin was recomputed from the resampled calibration engines' inner required margins (models were not refitted), held-out engines were re-evaluated at the new margins, and the paired difference Δ∗ in mean unused life over engines alerted in time by both models was recomputed. Intervals are the 2.5th and 97.5th percentiles of Δ∗; p = min{1, 2 min[P(Δ∗ ≤ 0), P(Δ∗ ≥ 0)]}, so its resolution is 0.001 and p = 0 is reported as p < 0.001."
)
Pp(
    "Multiplicity. The confirmatory family is the 15-cycle lead with the zero-miss target: RF − GB, GB − LR and CNN − GB in each subset, Holm-adjusted within the subset (three contrasts). Table S8 also gives Holm adjustment across all 12 contrasts. All other leads, the 2% target, the repeated partitions and the cap-150 analysis are exploratory; repeated partitions reuse the same engines and are not independent replications."
)
# ---------------- S1 accuracy
H("S1 Full Accuracy Results")
rows = [
    [
        "Subset",
        "Model",
        "MAE",
        "RMSE",
        "NASA",
        "Fold-avg. MAE / RMSE / NASA",
        "RMSE wins",
        "NASA wins",
        "Nested OOF RMSE",
    ]
]
groups = []
for s in SUBS:
    groups.append(len(rows))
    for m in ("LR", "Ridge", "RF", "GB", "CNN"):
        a = ACC[s][m]
        e = a["ens"]
        w = (f"{A[s]['wins']['RMSE'][m]:g}", f"{A[s]['wins']['NASA'][m]:g}") if m != "CNN" else ("–", "–")
        rows.append(
            [
                s if m == "LR" else "",
                m,
                f"{a['MAE'][0]:.2f} ± {a['MAE'][1]:.2f}",
                f"{a['RMSE'][0]:.2f} ± {a['RMSE'][1]:.2f}",
                f"{a['NASA'][0]:.0f} ± {a['NASA'][1]:.0f}",
                f"{e['MAE']:.2f} / {e['RMSE']:.2f} / {e['NASA']:.0f}",
                w[0],
                w[1],
                f"{a['oof_RMSE']:.2f}",
            ]
        )
cap(
    "Table S1.",
    "Official-test accuracy (capped labels) over the five fold models (mean ± SD), of the fold-averaged prediction, the number of folds (of five) in which each tabular model ranked first among the four tabular models (ties split), and the nested out-of-fold (OOF) RMSE over all training trajectories used by the decision lens.",
)
table(rows, [13, 12, 21, 21, 18, 37, 13, 13, 22], 7.5, groups)
# ---------------- S2 CNN
H("S2 Diagnostic of the CNN near Failure")
Pp(
    "Without a margin the CNN missed 97 to 100% of engines at a 15-cycle lead although it was the most accurate model out of fold on FD001 and FD003. Table S2 and Fig. S1 show why: its nested out-of-fold predictions have a floor near failure. Median predictions at failure were 16 to 25 cycles with narrow interquartile ranges on FD001, FD003 and FD004, so the predicted RUL rarely or never reaches 15 cycles while the engine still has 15 cycles left."
)
Pp(
    "Implementation checks. (i) Window alignment: each 30-cycle window ends at the current cycle; padding (repeating the first observation) affects only the first 29 cycles of an engine, far from failure. (ii) Rescaling: the network was trained on RUL/125 and its output multiplied by 125; predictions for healthy engines (true RUL ≥ 150, capped label 125) averaged 118 to 123 cycles. (iii) Bias pattern: a misalignment or rescaling error would shift or scale predictions throughout life; here the bias falls from about +16 to +23 cycles at failure to +5 to +12 at a true RUL of 30 to 50 (Table S2), with a small IQR at failure. (iv) In-sample check: refitting fold 0 of FD001 and predicting the training engines themselves gave the same floor (Table S3), so it is not an out-of-sample effect. (v) Ablation: removing the 0.5 dropout before the dense layer removed the floor (median 2.8 cycles at failure) with slightly better held-out fold RMSE, whereas training for 300 instead of 150 epochs raised it. The floor is therefore a property of this dropout configuration (as in Li et al.), not an implementation error, and the CNN's decision results hold for this configuration."
)
rows = [
    [
        "Subset",
        "Model",
        "Pred. at RUL 0: median [IQR]",
        "Pred. at RUL 15: median [IQR]",
        "Share ≤ 15 at RUL 15",
        "Pred. at RUL 30: median [IQR]",
        "Bias at 30 / 50",
        "Min pred.",
        "Healthy mean",
    ]
]
groups = []
for s in SUBS:
    groups.append(len(rows))
    for m in ("LR", "RF", "GB", "CNN"):
        d = DG[s]["dist"][m]
        q = lambda r: f"{f1(d[r]['med'])} [{f1(d[r]['q25'])}, {f1(d[r]['q75'])}]"
        rows.append(
            [
                s if m == "LR" else "",
                m,
                q("0"),
                q("15"),
                f"{d['15']['share_le_r']:.2f}",
                q("30"),
                f"{d['30']['mean_err']:+.1f} / {d['50']['mean_err']:+.1f}".replace("-", "−"),
                f1(DG[s]["healthy"][m]["min_pred"]),
                f1(DG[s]["healthy"][m]["mean_pred_rul_ge150"]),
            ]
        )
cap(
    "Table S2.",
    "Distribution of nested out-of-fold predicted RUL at selected true (uncapped) RUL values over all training engines. Bias: mean predicted minus true RUL. Min pred.: lowest prediction anywhere in the trajectories. Healthy mean: mean prediction at true RUL ≥ 150.",
)
table(rows, [13, 12, 26, 26, 13, 26, 20, 12, 14], 7.5, groups)
fig(
    "results/figures/figS1_cnn_floor.png",
    170,
    "Fig. S1.",
    "Median (line) and interquartile range (band) of nested out-of-fold predicted RUL against true RUL for linear regression, gradient boosting and the CNN. Dashed: perfect prediction; dotted: 15-cycle alarm level at m = 0.",
)
rows = [
    [
        "Run (FD001, folds 1–4 fitted)",
        "Engines",
        "Median pred. at RUL 0",
        "at RUL 15",
        "at RUL 30",
        "Min pred.",
        "RMSE",
    ]
]
for k, lab in (
    ("base", "As in paper (dropout 0.5, 150 epochs)"),
    ("dropout0", "No dropout, 150 epochs"),
    ("epochs300", "Dropout 0.5, 300 epochs"),
):
    for e, el in (("in", "training (in-sample)"), ("out", "held-out fold 0")):
        d = AB[k][e]
        rows.append(
            [
                lab if e == "in" else "",
                el,
                f1(d["0"][0]),
                f1(d["15"][0]),
                f1(d["30"][0]),
                f1(d["min"]),
                f"{d['rmse']:.2f}",
            ]
        )
cap(
    "Table S3.",
    "In-sample check and ablation for the CNN on FD001 (one fit on folds 1 to 4, seed 42). RMSE against capped labels over all cycles of the engines shown.",
)
table(rows, [52, 30, 22, 16, 16, 16, 14], 8.5, [3, 5])
# ---------------- S3 margins
H("S3 Why Linear Regression Wins at Long Lead Times")
Pp(
    "Table S4 decomposes the margins. Near failure (L = 10 or 15) the ensembles' out-of-fold errors at true RUL = L are tight (SD about 4 to 8 cycles) compared with linear regression (13 to 16), so they need smaller margins. At L = 30 their spread is similar to linear regression's, but the upper tail is heavier: the zero-miss margin, set by the hardest one or two calibration engines, exceeds the 95th-percentile requirement by 16 to 42 cycles against 13 to 18 for linear regression, and more calibration engines are outlying. On FD001, FD002 and FD004 the resulting alarm level L + m (91 to 109 cycles) lies below, and close to, the ensembles' predictions for healthy engines (115 to 119 cycles; Table S2), so alarms fire as soon as predictions leave that plateau, at a median true RUL of 94 to 112 cycles against 74 to 81 for linear regression. Unused life is computed as the true RUL at the alarm minus L, that is, assuming removal L cycles after the alarm. At the 95th-percentile margin (last column) the linear advantage shrinks on FD002 and reverses on FD004, but not on FD001, so the extreme calibration engines drive much, though not all, of the linear advantage at L = 30."
)
rows = [
    [
        "L",
        "Subset",
        "Model",
        "Zero-miss margins (folds)",
        "P95 req. (mean)",
        "Tail gap",
        "Outlying engines",
        "Error SD at RUL = L",
        "Median alarm RUL",
        "Unused, zero-miss",
        "Unused at P95 margin",
    ]
]
groups = []
for L in ("10", "15", "30"):
    for s in SUBS:
        groups.append(len(rows))
        for m in ("LR", "RF", "GB", "CNN"):
            x = DG[s]["margin"][L][m]
            rows.append(
                [
                    L if (s == "FD001" and m == "LR") else "",
                    s if m == "LR" else "",
                    m,
                    f"{min(x['zero'])}–{max(x['zero'])}",
                    f1(np.mean(x["p95"])),
                    f1(np.mean(np.array(x["zero"]) - np.array(x["p95"]))),
                    f"{np.mean(x['outliers']):.1f}",
                    f1(x["err_at_L_sd"]),
                    f"{x['med_alarm_rul']:.0f}",
                    f1(x["unused"]),
                    f1(x["unused_p95margin"]),
                ]
            )
cap(
    "Table S4.",
    "Margin decomposition by lead time L. Zero-miss margins: range over the five outer folds. P95 req.: 95th percentile of the calibration engines' required margins (mean over folds). Tail gap: zero-miss margin minus P95 requirement (mean over folds). Outlying engines: calibration engines with required margin above median + 3 IQR (mean per fold). Error SD: SD of out-of-fold prediction error at true RUL = L. Median alarm RUL: median true RUL at the alarm for timely-alerted held-out engines. Unused at P95 margin: unused life if each fold used its P95 requirement as margin.",
)
table(rows, [7, 13, 12, 18, 14, 12, 15, 15, 14, 15, 15], 7.5, groups)
# ---------------- S4 curves
H("S4 Out-of-Fold Miss-Rate Curves")
Pp(
    "For each model, the margin was chosen on the pooled nested out-of-fold trajectories as the smallest that leaves at most a fraction q of engines missed, and the unused life of timely-alerted engines was recorded (in-sample with respect to the margin, so descriptive). At L = 15 the tabular models differed by at most 6.3 cycles at q = 5% but by up to 15.5 cycles at q = 0, and the step from 5% to zero misses accounted for 21 to 55% of their zero-miss unused life: models separate in the tail of the error distribution."
)
C = SP["curve"]
rows = [
    ["L", "Subset"]
    + [f"{m} 5% / 0%" for m in ("LR", "Ridge", "RF", "GB", "CNN")]
    + ["Tabular spread 5% / 0%", "Step share, tabular"]
]
groups = []
for L in ("10", "15", "30"):
    groups.append(len(rows))
    for s in SUBS:
        c = C[s][L]
        tab = ("LR", "Ridge", "RF", "GB")
        u5 = [c[m]["0.05"]["unused"] for m in tab]
        u0 = [c[m]["0.0"]["unused"] for m in tab]
        sh = [100 * (c[m]["0.0"]["unused"] - c[m]["0.05"]["unused"]) / c[m]["0.0"]["unused"] for m in tab]
        rows.append(
            [L if s == "FD001" else "", s]
            + [
                f"{c[m]['0.05']['unused']:.1f} / {c[m]['0.0']['unused']:.1f}"
                for m in ("LR", "Ridge", "RF", "GB", "CNN")
            ]
            + [f"{max(u5) - min(u5):.1f} / {max(u0) - min(u0):.1f}", f"{min(sh):.0f}–{max(sh):.0f}%"]
        )
cap(
    "Table S5.",
    "Unused life on the pooled nested out-of-fold trajectories at 5% and 0% allowed misses, spread between the four tabular models, and the share of zero-miss unused life due to the step from 5% to zero misses (range over tabular models).",
)
table(rows, [7, 13, 19, 19, 19, 19, 19, 21, 17], 7.5, groups)
fig(
    "results/figures/figS2_miss_rate_curves.png",
    165,
    "Fig. S2.",
    "Unused life against allowed miss rate on the pooled nested out-of-fold trajectories, by subset (columns) and lead time (rows).",
)
# ---------------- S5 counts
H("S5 Significant Contrasts across Partitions")
cnt = SP["counts"]
rows = [["Contrast type", "Lead L", "Favoured", "Nested", "P1", "P2", "Total (of 48 or 24 cells)"]]
for key, lab in (
    ("linens|10|ensemble", ("Linear vs ensemble", "10", "ensemble")),
    ("linens|15|ensemble", ("", "15", "ensemble")),
    ("linens|30|linear", ("", "30", "linear")),
):
    c = cnt.get(key, {})
    rows.append(list(lab) + [str(c.get(p, 0)) for p in ("nested", "P1", "P2")] + [f"{sum(c.values())} of 48"])
for L in ("10", "15", "30"):
    for who in ("ensemble", "linear"):
        k = f"linens|{L}|{who}"
        if k in cnt or True:
            pass
opp = sum(
    v
    for k, v in ((k, sum(d.values())) for k, d in cnt.items())
    if k in ("linens|10|linear", "linens|15|linear", "linens|30|ensemble")
)
rows.append(["", "any", "opposite direction", "–", "–", "–", str(opp)])
for key, lab in (
    ("ens|10|RF", ("RF vs GB", "10", "RF")),
    ("ens|15|RF", ("", "15", "RF")),
    ("ens|30|GB", ("", "30", "GB")),
):
    c = cnt.get(key, {})
    rows.append(list(lab) + [str(c.get(p, 0)) for p in ("nested", "P1", "P2")] + [f"{sum(c.values())} of 24"])
opp2 = sum(sum(d.values()) for k, d in cnt.items() if k in ("ens|10|GB", "ens|15|GB", "ens|30|RF"))
rows.append(["", "any", "opposite direction", "–", "–", "–", str(opp2)])
cap(
    "Table S6.",
    "Counts of significant paired differences in unused life (95% interval excluding zero) by partition. Linear vs ensemble: GB − LR and RF − LR, two targets and four subsets per lead (16 cells per partition, 48 in total). RF vs GB: eight cells per lead and partition (24 in total). The partitions reuse the same engines, so these counts are not independent replications.",
)
table(rows, [30, 14, 26, 15, 12, 12, 36], 8.5, [4, 5, 8])
# ---------------- S6 cost
H("S6 Sensitivity to a Penalty per Missed Engine")
rows = [
    [
        "Pairs",
        "c (cycles per miss)",
        "Significant, kept",
        "Lost significance",
        "Reversed",
        "Newly significant",
    ]
]


def costcount(sel, c):
    k = l = r = n = 0
    for s in SUBS:
        for L in ("10", "15", "30"):
            for a in ("0.0", "0.02"):
                for pair in sel:
                    v = D[s][L][a]["pairs"][pair]
                    lo, hi = v["cost"][c]["ci95"]
                    sc = 0 if lo <= 0 <= hi else int(np.sign(v["cost"][c]["mean"]))
                    if v["sig"] and sc == v["sig"]:
                        k += 1
                    elif v["sig"] and sc == 0:
                        l += 1
                    elif v["sig"] and sc == -v["sig"]:
                        r += 1
                    elif not v["sig"] and sc:
                        n += 1
    return k, l, r, n


for sel, lab in (
    (("RF-GB", "GB-LR", "RF-LR"), "Tabular (RF − GB, GB − LR, RF − LR)"),
    (("CNN-GB", "CNN-LR", "CNN-RF"), "CNN vs tabular"),
):
    for c in ("25", "100", "400"):
        rows.append([lab if c == "25" else "", c] + [str(x) for x in costcount(sel, c)])
cap(
    "Table S7.",
    "Nested decision lens over all 24 cells (3 leads × 2 targets × 4 subsets): comparison of the cost per engine (unused life of timely-alerted engines plus c per missed engine, divided by all engines) with the paired unused-life comparison. At c = 400, 16 of the 17 significant tabular ensemble advantages at 10 and 15 cycles lost significance, while all 7 significant linear advantages at 30 cycles remained.",
)
table(rows, [62, 24, 22, 22, 18, 22], 8.5, [4])
# ---------------- S7 Holm
H("S7 Primary Contrasts and Multiplicity")


def holm(ps):
    o = np.argsort(ps)
    m = len(ps)
    adj = np.empty(m)
    run = 0
    for rnk, i in enumerate(o):
        run = max(run, min(1, (m - rnk) * ps[i]))
        adj[i] = run
    return adj


lab = []
ps = []
within = []
for s in SUBS:
    pr = D[s]["15"]["0.0"]["pairs"]
    sub = []
    for p in ("RF-GB", "GB-LR", "CNN-GB"):
        lab.append((s, p))
        ps.append(max(pr[p]["p"], 0.0005))
        sub.append(ps[-1])
    within += list(holm(np.array(sub)))
across = holm(np.array(ps))
fp = lambda x: "< 0.001" if x < 0.001 else ("< 0.002" if x < 0.002 else f"{x:.3f}")
rows = [["Subset", "Contrast", "Δ unused [95% CI]", "n", "p", "Holm within subset", "Holm across 12"]]
for (s, p), raw, w, a in zip(lab, ps, within, across):
    v = D[s]["15"]["0.0"]["pairs"][p]
    rows.append(
        [
            s if p == "RF-GB" else "",
            p.replace("-", " − "),
            f"{f1(v['mean'])} [{f1(v['ci95'][0])}, {f1(v['ci95'][1])}]",
            str(v["n"]),
            "< 0.001" if v["p"] == 0 else f"{v['p']:.3f}",
            "< 0.003" if v["p"] == 0 else f"{w:.3f}",
            "< 0.012" if v["p"] == 0 else f"{a:.3f}",
        ]
    )
cap(
    "Table S8.",
    "Confirmatory family: 15-cycle lead, zero-miss target. Negative Δ favours the first model. With p < 0.001 (no replicate crossing zero), adjusted values are bounds.",
)
table(rows, [16, 22, 36, 12, 16, 26, 24], 8.5, [4, 7, 10])
# ---------------- S8 full decision table (landscape)
from docx.enum.section import WD_SECTION

ns = doc.add_section(WD_SECTION.NEW_PAGE)
ns.orientation = WD_ORIENT.LANDSCAPE
ns.page_width, ns.page_height = Mm(297), Mm(210)
for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(ns, a, Mm(18))
H("S8 Full Nested Decision Results")
rows = [["L", "Target", "Subset", "LR", "Ridge", "RF", "GB", "CNN", "Fixed age"]]
groups = []
for L in ("10", "15", "30"):
    for a, al in (("0.0", "0%"), ("0.02", "2%")):
        groups.append(len(rows))
        for s in SUBS:
            md = D[s][L][a]["models"]
            tm = D[s][L]["time"]
            rows.append(
                [L if (a == "0.0" and s == "FD001") else "", al if s == "FD001" else "", s]
                + [
                    f"{md[m]['unused']:.1f} ({md[m]['misses']}) m{min(md[m]['margins'])}–{max(md[m]['margins'])}"
                    for m in ("LR", "Ridge", "RF", "GB", "CNN")
                ]
                + [f"{tm['unused']:.1f} ({tm['misses']})" if a == "0.0" else "–"]
            )
cap(
    "Table S9.",
    "Unused life of timely-alerted engines (held-out misses) and fold margins for every lead time and calibration miss target under nested calibration. Engines per subset: 100, 260, 100 and 249. Fixed age: zero-miss fixed-age replacement calibrated on the same folds.",
)
table(rows, [9, 13, 14, 38, 38, 38, 38, 38, 25], 8, groups)
rows = [["L", "Target", "Subset", "RF − GB", "GB − LR", "RF − LR", "CNN − GB", "CNN − LR"]]
groups = []
for L in ("10", "15", "30"):
    for a, al in (("0.0", "0%"), ("0.02", "2%")):
        groups.append(len(rows))
        for s in SUBS:
            pr = D[s][L][a]["pairs"]
            rows.append(
                [L if (a == "0.0" and s == "FD001") else "", al if s == "FD001" else "", s]
                + [
                    f"{f1(pr[k]['mean'])} [{f1(pr[k]['ci95'][0])}, {f1(pr[k]['ci95'][1])}]"
                    + ("*" if pr[k]["sig"] else "")
                    for k in ("RF-GB", "GB-LR", "RF-LR", "CNN-GB", "CNN-LR")
                ]
            )
doc.add_page_break()
cap(
    "Table S10.",
    "Paired differences in unused life over engines alerted in time by both models, with 95% bootstrap intervals including margin re-estimation (* interval excludes zero). Negative values favour the first model. Only the 15-cycle, 0% rows are confirmatory.",
)
table(rows, [9, 13, 14, 41, 41, 41, 41, 41], 8, groups)
Pp(
    "Code, model predictions and all intermediate results: https://github.com/toochii2/isdia2027-rul-decision-evaluation",
    it=True,
    size=9.5,
)
doc.save("docs/supplementary_material.docx")
print("saved")
